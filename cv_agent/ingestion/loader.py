"""
CV ingestion module.

Responsibilities:
  - Detect file type (PDF / DOCX).
  - Load and parse the document.
  - Split into overlapping chunks suitable for embedding.
  - Enrich each chunk's metadata with: source, page, chunk_id, section.
"""

from __future__ import annotations

import logging
import time
from pathlib import Path

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from cv_agent.config import settings

logger = logging.getLogger(__name__)


# ── Section inference ─────────────────────────────────────────────────────────

# Ordered list of (section_label, keyword_list) pairs.
# The first match wins; order matters (more specific → more general).
_SECTION_PATTERNS: list[tuple[str, list[str]]] = [
    ("Contact", ["email", "phone", "address", "linkedin", "github", "twitter", "contact"]),
    ("Summary", ["summary", "profile", "objective", "about me", "overview"]),
    ("Experience", ["experience", "employment", "work history", "career history", "positions held"]),
    ("Education", ["education", "academic", "qualification", "degree", "university", "college", "school"]),
    ("Skills", ["skills", "technologies", "tools", "competencies", "proficiencies", "technical skills"]),
    ("Certifications", ["certification", "certificate", "license", "accreditation", "credential"]),
    ("Projects", ["project", "portfolio", "open source", "side project"]),
    ("Publications", ["publication", "paper", "research", "journal", "conference"]),
    ("Awards", ["award", "achievement", "honor", "recognition", "accomplishment"]),
    ("Languages", ["languages", "spoken languages", "language proficiency"]),
    ("Interests", ["interest", "hobbies", "activities", "volunteer"]),
]


def _infer_section(text: str) -> str:
    """
    Attempt to infer the CV section from a chunk's text content.

    Scans the first 300 characters (where section headings typically appear)
    for known keyword patterns. Returns "General" if no match is found.
    """
    scan_text = text[:300].lower()
    for section_label, keywords in _SECTION_PATTERNS:
        if any(kw in scan_text for kw in keywords):
            return section_label
    return "General"


# ── Loaders ───────────────────────────────────────────────────────────────────


def _load_pdf(path: Path) -> list[Document]:
    """Load a PDF file using pypdf via LangChain's PyPDFLoader."""
    try:
        from langchain_community.document_loaders import PyPDFLoader
    except ImportError as exc:
        raise ImportError(
            "pypdf is required for PDF ingestion. "
            "Install it with: pip install pypdf"
        ) from exc

    logger.info("INGESTION | Loading CV: %s", path)
    loader = PyPDFLoader(str(path))
    docs = loader.load()
    total_chars = sum(len(d.page_content) for d in docs)
    logger.info(
        "INGESTION | Extracted %d characters across %d page(s) from PDF",
        total_chars,
        len(docs),
    )
    return docs


def _load_docx(path: Path) -> list[Document]:
    """Load a DOCX file using python-docx."""
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise ImportError(
            "python-docx is required for DOCX ingestion. "
            "Install it with: pip install python-docx"
        ) from exc

    logger.info("INGESTION | Loading CV: %s", path)
    document = docx.Document(str(path))
    full_text = "\n".join(
        para.text for para in document.paragraphs if para.text.strip()
    )
    if not full_text.strip():
        raise ValueError(f"DOCX file appears to be empty: {path}")

    docs = [Document(page_content=full_text, metadata={"source": str(path), "page": 0})]
    logger.info("INGESTION | Extracted %d characters from DOCX", len(full_text))
    return docs


def _extract_candidate_name(cv_path: Path, raw_docs: list[Document]) -> str:
    """Infer candidate name from first text line or filename."""
    stem = cv_path.stem.replace("_", " ").replace("-", " ")
    cleaned_stem = " ".join([w for w in stem.split() if w.lower() not in ["resume", "cv", "pdf", "docx"]])

    if raw_docs and raw_docs[0].page_content.strip():
        first_line = raw_docs[0].page_content.strip().split("\n")[0].strip()
        bad_words = [
            "resume", "curriculum", "page", "mba", "bba", "b.tech", "btech",
            "engineer", "profile", "summary", "experience", "education", "skills",
            "candidate", "contact", "email", "phone"
        ]
        if 2 <= len(first_line) <= 40 and not any(k in first_line.lower() for k in bad_words):
            return first_line

    return cleaned_stem.title() if cleaned_stem else stem.title()


# ── Public API ────────────────────────────────────────────────────────────────


def load_cv(path: Path | str | None = None) -> list[Document]:
    """
    Load and chunk a single CV from *path* (defaults to settings.cv_file_path).

    Each returned chunk carries full metadata:
        source      — file path string
        filename    — basename of file (e.g. Hari_Resume.pdf)
        candidate   — candidate name
        page        — 0-based page index (from PDF; 0 for DOCX)
        chunk_id    — sequential integer index across all chunks
        section     — inferred CV section label (e.g. "Experience", "Skills")

    Returns:
        List of chunked LangChain Documents ready for embedding.
    """
    cv_path = Path(path or settings.cv_file_path)

    # ── Validate ──────────────────────────────────────────────────────────────
    if not cv_path.exists():
        raise FileNotFoundError(
            f"CV file not found at '{cv_path}'. "
            "Please place your CV (PDF or DOCX) at the configured path and retry."
        )

    suffix = cv_path.suffix.lower()
    if suffix not in {".pdf", ".docx"}:
        raise ValueError(
            f"Unsupported file type '{suffix}'. Only .pdf and .docx are supported."
        )

    # ── Load ──────────────────────────────────────────────────────────────────
    t0 = time.perf_counter()
    if suffix == ".pdf":
        raw_docs = _load_pdf(cv_path)
    else:
        raw_docs = _load_docx(cv_path)

    if not raw_docs:
        raise ValueError(f"No content could be extracted from the CV file: {cv_path}")

    candidate_name = _extract_candidate_name(cv_path, raw_docs)

    # ── Chunk ─────────────────────────────────────────────────────────────────
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        separators=["\n\n", "\n", " ", ""],
    )
    chunks = splitter.split_documents(raw_docs)

    if not chunks:
        raise ValueError(
            f"Chunking produced zero chunks for '{cv_path.name}'. "
            "Check that the CV file has readable text content."
        )

    logger.info(
        "CHUNKING | '%s' (Candidate: %s) -> %d chunks | chunk_size=%d | overlap=%d",
        cv_path.name,
        candidate_name,
        len(chunks),
        settings.chunk_size,
        settings.chunk_overlap,
    )

    # ── Enrich metadata ───────────────────────────────────────────────────────
    for chunk_id, chunk in enumerate(chunks):
        chunk.metadata["chunk_id"] = chunk_id
        chunk.metadata["section"] = _infer_section(chunk.page_content)
        chunk.metadata["source"] = str(cv_path)
        chunk.metadata["filename"] = cv_path.name
        chunk.metadata["candidate"] = candidate_name

    elapsed = time.perf_counter() - t0
    logger.info(
        "INGESTION | Complete | %s -> %d chunk(s) | elapsed=%.2fs",
        cv_path.name,
        len(chunks),
        elapsed,
    )

    return chunks


def load_multiple_cvs(paths: list[Path | str] | Path | str | None = None) -> list[Document]:
    """
    Load and chunk multiple CVs from a list of paths, a single path, or all
    PDF/DOCX files in the data directory.

    Returns:
        Combined list of chunked Documents from all specified CVs.
    """
    if paths is None:
        # Default: check settings.cv_file_path or scan data/
        if settings.cv_file_path.exists():
            file_paths = [settings.cv_file_path]
        else:
            data_dir = Path("data")
            file_paths = list(data_dir.glob("*.pdf")) + list(data_dir.glob("*.docx"))
    elif isinstance(paths, (list, tuple)):
        file_paths = [Path(p) for p in paths]
    elif Path(paths).is_dir():
        dir_path = Path(paths)
        file_paths = list(dir_path.glob("*.pdf")) + list(dir_path.glob("*.docx"))
    else:
        file_paths = [Path(paths)]

    if not file_paths:
        raise FileNotFoundError("No CV files found to ingest (.pdf or .docx).")

    all_chunks: list[Document] = []
    global_id = 0

    for fp in file_paths:
        if not fp.exists():
            logger.warning("INGESTION | Skipping non-existent path: %s", fp)
            continue
        try:
            doc_chunks = load_cv(fp)
            for c in doc_chunks:
                c.metadata["global_chunk_id"] = global_id
                global_id += 1
            all_chunks.extend(doc_chunks)
        except Exception as exc:
            logger.error("INGESTION | Error loading %s: %s", fp, exc)

    logger.info(
        "INGESTION | Total multi-CV ingestion: %d file(s) -> %d total chunk(s)",
        len(file_paths),
        len(all_chunks),
    )
    return all_chunks
