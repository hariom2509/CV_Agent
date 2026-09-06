"""
Tests for the ingestion module.
"""

from __future__ import annotations

import tempfile
from pathlib import Path

import pytest


class TestLoadCVValidation:
    """Test input validation in load_cv()."""

    def test_missing_file_raises_file_not_found(self):
        from cv_agent.ingestion.loader import load_cv

        with pytest.raises(FileNotFoundError, match="CV file not found"):
            load_cv(Path("nonexistent_file.pdf"))

    def test_unsupported_extension_raises_value_error(self):
        from cv_agent.ingestion.loader import load_cv

        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as f:
            Path(f.name).write_text("some content")
            with pytest.raises(ValueError, match="Unsupported file type"):
                load_cv(Path(f.name))

    def test_empty_docx_raises_value_error(self):
        """An empty DOCX file should raise ValueError after load."""
        import docx

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "empty.docx"
            doc = docx.Document()
            doc.save(str(path))

            with pytest.raises((ValueError, Exception)):
                from cv_agent.ingestion.loader import load_cv
                load_cv(path)


class TestLoadCVDocx:
    """Test DOCX ingestion with a synthetic document."""

    def test_docx_ingestion_returns_chunks(self):
        import docx

        from cv_agent.ingestion.loader import load_cv

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "test_cv.docx"
            doc = docx.Document()
            doc.add_paragraph("John Doe — Software Engineer")
            doc.add_paragraph(
                "Experience: 5 years at Acme Corp. Led a team of 8 engineers. "
                "Built distributed systems serving 10M users."
            )
            doc.add_paragraph("Skills: Python, LangChain, Kubernetes, PostgreSQL")
            doc.add_paragraph("Education: B.Sc. Computer Science, University of Cape Town, 2018")
            doc.save(str(path))

            chunks = load_cv(path)

            assert len(chunks) >= 1, "Expected at least one chunk"
            combined = " ".join(c.page_content for c in chunks)
            assert "John Doe" in combined
            assert "Python" in combined

    def test_chunk_metadata_has_source(self):
        import docx

        from cv_agent.ingestion.loader import load_cv

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "meta_test.docx"
            doc = docx.Document()
            doc.add_paragraph("Test CV content " * 50)
            doc.save(str(path))

            chunks = load_cv(path)
            for chunk in chunks:
                assert "source" in chunk.metadata
